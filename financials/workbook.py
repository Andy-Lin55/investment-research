"""Portable review workbooks for the opt-in cloud runner, never production imports.

Formula caches are checked Python projections, not values read from edited Excel.
Only explicit assumption cells may be imported; all outputs are rebuilt in Python.
"""
from copy import deepcopy
from datetime import datetime
from io import BytesIO
from zipfile import ZipFile

from financials.forecast import BASE_FIELDS, DRIVERS, number, project

CASES = ('base', 'bull', 'bear')
ROWS = {5:'收入',6:'营业成本',7:'经营费用',8:'折旧',9:'利息',10:'税前利润',11:'所得税',
        12:'净利润',13:'归母净利润',14:'每股盈利',17:'现金',18:'应收账款',19:'存货',20:'固定资产',
        21:'其他资产',22:'应付账款',23:'债务',24:'其他负债',25:'权益',26:'股数',
        29:'经营现金流',30:'投资现金流',31:'融资现金流',32:'汇兑现金影响',33:'资本开支',
        34:'股息',35:'现金变化',38:'资产负债差额',39:'现金勾稽差额',40:'资金缺口',43:'估值市盈率',44:'情景每股估值'}
BASE_ROWS = dict(zip(BASE_FIELDS, (5,17,18,19,20,21,22,23,24,25,26)))


def checked(model):
    rebuilt = project(model['baseline'], model['assumptions'], as_of=model['as_of'])
    for k in ('model_version','issuer_id','currency','status','scenarios'):
        if rebuilt[k] != model[k]:raise ValueError('model snapshot failed reproducibility check')
    if sorted(rebuilt['missing'])!=sorted(model['missing']):raise ValueError('model missing inputs differ')
    if len(model['assumptions']['base']) > 10:raise ValueError('review at most ten annual forecast periods')
    return rebuilt


def export_model(model):
    checked(model)
    # This optional dependency belongs only to requirements-research.txt.
    import xlsxwriter
    from xlsxwriter.utility import xl_col_to_name
    output = BytesIO()
    book = xlsxwriter.Workbook(output, {'in_memory':True,'strings_to_formulas':False,'strings_to_urls':False})
    book.set_properties({'title':'持仓研究三表模型｜待人工审核',
                         'created':datetime.fromisoformat(model['as_of'])})
    book.set_calc_mode('auto')
    header = book.add_format({'bold':True,'bg_color':'#183F37','font_color':'white','text_wrap':True})
    label = book.add_format({'text_wrap':True,'valign':'top'})
    numeric = book.add_format({'num_format':'#,##0.00;[Red](#,##0.00);–'})
    editable = book.add_format({'font_color':'#0000FF','num_format':'0.0000','bg_color':'#FFF4CD'})
    sheets = {name:book.add_worksheet(name) for name in ('Review','Baseline','Assumptions','Base','Bull','Bear','Changes')}
    for sheet in sheets.values():
        sheet.set_column('A:A',28);sheet.set_column('B:D',24);sheet.set_column('E:E',65)
        sheet.freeze_panes(1,1);sheet.set_landscape();sheet.fit_to_pages(1,0)
    review = sheets['Review']
    review.write_row(0,0,['持仓研究三表与情景估值','待人工核验'],header)
    metadata = [('Issuer',model['issuer_id']),('Model version',model['model_version']),('Status',model['status']),
                ('As of',model['as_of']),('Currency',model['currency']),('Baseline revision',model['baseline_revision']),
                ('Period end',model['baseline']['period_end']),('Source',model['baseline']['source_locator']),
                ('Cash scope confirmed',model['baseline']['cash_scope_confirmed'])]
    for i,(name,value) in enumerate(metadata,1):review.write_row(i,0,[name,value],label)
    review.set_column('B:B',90)
    for i,text in enumerate([
        '金额为原始货币单位，股数为未缩放股数；仅适用已校验的工业企业年度模型。',
        'Base / Bull / Bear 分别为显式假设情景，不代表市场预测；未填写的假设不按零计算。',
        '只编辑 Assumptions 的 Value 与 Source；导入时核对原模型版本，重新计算全部输出，仍须人工批准。',
        '报表展示包含 Python 计算缓存；Excel 打开后自动重算，导出不代表原生 Excel 重算已经验收。',
        '未解决的假设：'+', '.join(model['missing'])],12):
        review.write(i,0,text,label);review.merge_range(i,0,i,1,text,label);review.set_row(i,34)
    baseline = sheets['Baseline'];baseline.write_row(0,0,['Field','Value','Source'],header)
    for i,field in enumerate(BASE_FIELDS,1):
        baseline.write_row(i,0,[field,model['baseline'][field],model['baseline']['source_locator']],label)
    baseline.set_column('C:C',65)
    assumptions = sheets['Assumptions'];assumptions.write_row(0,0,['Scenario','Year','Driver','Value','Source','Input valid','All inputs ready'],header)
    assumptions.set_column('F:G',22)
    refs={};index=1
    for name in CASES:
        for period in model['assumptions'][name]:
            for field in DRIVERS:
                assumptions.write_row(index,0,[name,period['year'],field],label)
                value=period.get('values',{}).get(field)
                if number(value):assumptions.write_number(index,3,value,editable)
                else:assumptions.write_blank(index,3,None,editable)
                assumptions.write_string(index,4,period.get('sources',{}).get(field) or '',label)
                source=period.get('sources',{}).get(field)
                valid=number(value) and isinstance(source,str) and bool(source.strip())
                assumptions.write_formula(index,5,f'=IF(ISNUMBER(D{index+1}),LEN(TRIM(E{index+1}))>0,FALSE)',None,valid)
                assumptions.set_row(index,30)
                refs[name,period['year'],field]=f'Assumptions!D{index+1}';index+=1
    assumptions.autofilter(0,0,index-1,4)
    assumptions.write_formula(1,6,f'=COUNTIF(F2:F{index},TRUE)={index-1}',None,not bool(model['missing']))
    for name in CASES:
        sheet=sheets[name.title()];sheet.write_row(0,0,[name.upper()+'｜预测假设','历史基准'],header)
        for row,text in ROWS.items():sheet.write(row-1,0,text,label)
        for i,field in enumerate(BASE_FIELDS,2):
            sheet.write_formula(BASE_ROWS[field]-1,1,f'=Baseline!B{i}',numeric,model['baseline'][field])
        for col,period in enumerate(model['assumptions'][name],2):
            year=period['year'];c=xl_col_to_name(col);p=xl_col_to_name(col-1)
            sheet.write(0,col,str(year)+'E',header)
            # Keep formulas in pending templates so filling every required input
            # can recalculate. Clearing any required value/source suppresses all
            # scenarios instead of treating a missing assumption as zero.
            a={k:refs[name,year,k] for k in DRIVERS}
            f={5:f'{p}5*(1+{a["growth"]})',6:f'{c}5*{a["cogs_ratio"]}',7:f'{c}5*{a["opex_ratio"]}',
               8:f'{p}20*{a["dep_rate"]}',9:f'({p}23+{c}23)/2*{a["interest_rate"]}',
               10:f'{c}5-{c}6-{c}7-{c}8-{c}9',11:f'MAX({c}10,0)*{a["tax_rate"]}',12:f'{c}10-{c}11',
               13:f'{c}12*(1-{a["minority_ratio"]})',14:f'{c}13/(({p}26+{c}26)/2)',
               17:f'{p}17+{c}29+{c}30+{c}31+{c}32',18:f'{c}5*{a["receivable_days"]}/365',
               19:f'{c}6*{a["inventory_days"]}/365',20:f'{p}20+{c}33-{c}8',21:f'{p}21',
               22:f'{c}6*{a["payable_days"]}/365',23:f'{p}23+{a["debt_change"]}',24:f'{p}24',
               25:f'{p}25+{c}12-{c}34+{a["equity_issuance"]}+{a["other_equity_change"]}',
               26:f'{p}26+{a["share_change"]}',
               29:f'{c}12+{c}8-(({c}18-{p}18)+({c}19-{p}19)-({c}22-{p}22))',
               30:f'-{c}33',31:f'{a["debt_change"]}+{a["equity_issuance"]}-{c}34',32:a['fx_cash'],
               33:f'{c}5*{a["capex_ratio"]}',34:f'MAX({c}12,0)*{a["dividend_payout"]}',35:f'{c}17-{p}17',
               38:f'SUM({c}17:{c}21)-SUM({c}22:{c}25)',39:f'{c}35-{c}29-{c}30-{c}31-{c}32',
               40:f'MAX(0,-{c}17)',43:a['pe_multiple'],44:f'IF({c}14>0,{c}14*{c}43,"")'}
            caches={r:'' for r in f}
            if model['scenarios']:
                row=model['scenarios'][name][col-2];inc=row['income'];bal=row['balance'];cf=row['cashflow']
                caches.update(dict(zip(range(5,15),[inc[k] for k in ('revenue','cogs','opex','depreciation','interest','pretax','tax','net_income','parent_net_income','eps')])))
                caches.update(dict(zip(range(17,27),[bal[k] for k in ('cash','receivables','inventory','ppe','other_assets','payables','debt','other_liabilities','equity','shares')])))
                caches.update(dict(zip(range(29,36),[cf[k] for k in ('operating','investing','financing','fx','capex','dividends','change')])))
                caches.update({38:row['checks']['balance_difference'],39:row['checks']['cash_difference'],
                               40:row['funding_shortfall'],43:row['valuation']['multiple'],44:row['valuation']['price']})
            for r,formula in f.items():
                sheet.write_formula(r-1,col,'=IF(Assumptions!$G$2,'+formula+',"")',numeric,caches[r] if caches[r] is not None else '')
        last=xl_col_to_name(len(model['assumptions'][name])+1)
        sheet.conditional_format(f'C38:{last}39',{'type':'cell','criteria':'not between','minimum':-1e-6,'maximum':1e-6,
            'format':book.add_format({'bg_color':'#FFDAD6'})})
    changes=sheets['Changes'];changes.write_row(0,0,['Scenario','Year','Driver','Previous','Current','Source'],header)
    changes.set_column('F:F',65)
    for i,change in enumerate(model.get('assumption_changes',[]),1):
        changes.write_row(i,0,[change[k] for k in ('scenario','year','field','previous','current','source')],label)
    book.close();return output.getvalue()


def import_assumptions(data,original):
    checked(original)
    from openpyxl import load_workbook
    if len(data)>10_000_000:raise ValueError('bounded workbook required')
    with ZipFile(BytesIO(data)) as zipped:
        if (sum(i.file_size for i in zipped.infolist())>30_000_000 or len(zipped.infolist())>100
                or any('externalLinks/' in i.filename or 'vbaProject' in i.filename for i in zipped.infolist())):
            raise ValueError('unsupported or oversized workbook')
    book=load_workbook(BytesIO(data),data_only=False,keep_links=False)
    if book.sheetnames!=['Review','Baseline','Assumptions','Base','Bull','Bear','Changes']:
        raise ValueError('expected review workbook layout')
    identities={2:original['issuer_id'],3:original['model_version'],5:original['as_of'],6:original['currency'],
                7:original['baseline_revision'],8:original['baseline']['period_end'],
                9:original['baseline']['source_locator'],10:original['baseline']['cash_scope_confirmed']}
    if any(book['Review'].cell(r,2).value!=v for r,v in identities.items()):
        raise ValueError('workbook identity or baseline metadata changed')
    for i,key in enumerate(BASE_FIELDS,2):
        if [book['Baseline'].cell(i,c).value for c in (1,2,3)]!=[key,original['baseline'][key],original['baseline']['source_locator']]:
            raise ValueError('baseline changed; supply a separately verified source baseline')
    s=deepcopy(original['assumptions']);expected=[(name,row['year'],field) for name in CASES for row in s[name] for field in DRIVERS]
    sheet=book['Assumptions']
    if sheet.max_row!=len(expected)+1:raise ValueError('assumption row count differs')
    positions={(name,row['year']):row for name in CASES for row in s[name]}
    for i,(name,year,field) in enumerate(expected,2):
        if tuple(sheet.cell(i,c).value for c in (1,2,3))!=(name,year,field):raise ValueError('assumption identity differs')
        value,source=(sheet.cell(i,c) for c in (4,5))
        if value.data_type=='f' or source.data_type=='f':raise ValueError('literal assumptions and sources required')
        if value.value is not None and (not number(value.value) or not isinstance(source.value,str) or not source.value.strip()):
            raise ValueError('finite numeric assumption with explicit source required')
        row=positions[name,year]
        previous=row.get('values',{}).get(field)
        # Excel normalizes integral doubles (0.0 -> 0). Preserve the authoritative
        # original representation for equal numbers, so a save is not a revision.
        imported=previous if number(previous) and number(value.value) and previous==value.value else value.value
        row.setdefault('values',{})[field]=imported;row.setdefault('sources',{})[field]=source.value or ''
    return project(original['baseline'],s,as_of=original['as_of'],previous=original)
