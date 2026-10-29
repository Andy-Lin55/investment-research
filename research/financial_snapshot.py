"""Read a specifically mapped historical workbook, not arbitrary human forecasts."""
from datetime import date
from io import BytesIO
from zipfile import ZipFile
import hashlib,re
from financials.forecast import number

SHEETS=('利润表','资产负债表','现金流量表')


def from_workbook(data,*,issuer_id,source,as_of,expected_company):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import get_column_letter
    if not issuer_id or not expected_company:raise ValueError('confirmed issuer and company title required')
    if source.get('sha256')!=hashlib.sha256(data).hexdigest():raise ValueError('workbook checksum differs')
    date.fromisoformat(as_of)
    if len(data)>40_000_000:raise ValueError('bounded workbook required')
    with ZipFile(BytesIO(data)) as z:
        if sum(i.file_size for i in z.infolist())>100_000_000:raise ValueError('expanded workbook too large')
        if any('externalLinks/' in i.filename or 'vbaProject' in i.filename for i in z.infolist()):
            raise ValueError('external-linked/macro workbook needs separate review')
    book=load_workbook(BytesIO(data),data_only=True,read_only=True,keep_links=False)
    tables=[];gaps=[]
    try:
        for name in SHEETS:
            s=book[name]
            if s.max_row is None or s.max_column is None:s.calculate_dimension(force=True)
            if s.max_row>200 or s.max_column>25:raise ValueError('bounded historical table required')
            grid=list(s.iter_rows(min_row=1,max_row=s.max_row,max_col=s.max_column))
            def cell(r,c):return grid[r-1][c-1]
            if (not str(cell(2,2).value or '').startswith(expected_company+'｜')
                    or cell(5,2).value!='科目' or cell(5,3).value!='展示单位'):
                raise ValueError('explicit historical-workpaper template required')
            cols=[c for c in range(4,s.max_column+1) if cell(5,c).value is not None]
            headers=[str(cell(5,c).value) for c in cols]
            if not headers or len(headers)!=len(set(headers)) or any(not re.fullmatch(r'\d{4}(?:年度|-\d{2}-\d{2}(?:累计|期末))A',h) for h in headers):
                raise ValueError('explicit unique actual-period labels required')
            for h in headers:
                end=date.fromisoformat(h[:4]+'-12-31' if '年度A' in h else h[:10])
                if end>date.fromisoformat(as_of):raise ValueError('future financial period labelled actual')
            rows=[]
            for r in range(6,s.max_row+1):
                label=cell(r,2).value;unit=cell(r,3).value
                if label is None:continue
                if not isinstance(label,str) or unit not in {'亿元','元/股','亿股','%','倍','天'}:
                    raise ValueError('explicit supported display units required')
                values=[];cells=[]
                for c in cols:
                    source_cell=cell(r,c);value=source_cell.value;coordinate=get_column_letter(c)+str(r)
                    if value is not None and not number(value):
                        gaps.append(name+'!'+coordinate+':invalid_cached_value');value=None
                    if value is not None and unit=='%' and '%' in source_cell.number_format:value*=100
                    values.append(value);cells.append(name+'!'+coordinate)
                rows.append(dict(label=label,unit=unit,values=values,cells=cells))
            tables.append(dict(title=name,periods=headers,rows=rows))
    finally:book.close()
    return dict(schema_version=1,issuer_id=issuer_id,as_of=as_of,source=dict(source),tables=tables,gaps=gaps,
                human_review='pending',value_kind='historical_cached_workbook',
                note='Historical workbook display values only; periods/units remain explicit. No approved forecast inferred. 货币资金与现金等价物不混同。')
