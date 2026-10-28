"""Catch formula/cached-value divergence, missing-input fabrication and unsafe edits."""
from copy import deepcopy
from io import BytesIO
import importlib
import re
import pytest
from openpyxl import load_workbook
from financials.forecast import project
from tests.test_forecast_model import baseline, scenarios


def api():
    try:return importlib.import_module('financials.workbook')
    except ModuleNotFoundError:pytest.fail('Reviewable forecast workbook is not implemented')


def fixture():return project(baseline(),scenarios(),as_of='2026-10-27')


def test_workbook_has_linked_statements_and_independently_checked_results():
    data=api().export_model(fixture())
    formulas=load_workbook(BytesIO(data));values=load_workbook(BytesIO(data),data_only=True)
    assert formulas.sheetnames==['Review','Baseline','Assumptions','Base','Bull','Bear','Changes']
    assert formulas['Base']['C5'].value=='=IF(Assumptions!$G$2,B5*(1+Assumptions!D2),"")'
    assert values['Base']['C5'].value==1100
    assert values['Base']['C12'].value==217.5
    assert values['Base']['C25'].value==452.25
    assert values['Base']['C44'].value==pytest.approx(26.1)
    assert values['Base']['C38'].value==pytest.approx(0,abs=1e-8)
    assert values['Base']['C39'].value==pytest.approx(0,abs=1e-8)
    assert formulas['Assumptions']['E2'].value=='explicit synthetic assumption'
    assert api().export_model(fixture())==data


def test_later_year_links_prior_balance_and_does_not_reuse_opening_cash():
    s=scenarios()
    for rows in s.values():rows.append(dict(deepcopy(rows[0]),year=2027))
    m=project(baseline(),s,as_of='2026-10-27');w=load_workbook(BytesIO(api().export_model(m)))
    assert ',C5*' in w['Base']['D5'].value
    assert ',C17+' in w['Base']['D17'].value
    cached=load_workbook(BytesIO(api().export_model(m)),data_only=True)
    assert cached['Base']['D5'].value==pytest.approx(1210)
    assert cached['Base']['D38'].value==pytest.approx(0,abs=1e-8)


def test_missing_assumption_produces_editable_blank_without_valuation():
    s=scenarios();s['base'][0]['values']['growth']=None
    m=project(baseline(),s,as_of='2026-10-27');w=load_workbook(BytesIO(api().export_model(m)))
    assert w['Assumptions']['D2'].value is None
    assert load_workbook(BytesIO(api().export_model(m)),data_only=True)['Base']['C44'].value is None
    assert 'pending_assumptions' in w['Review']['B4'].value


def test_only_explicit_human_assumption_edits_can_be_imported():
    m=fixture();w=load_workbook(BytesIO(api().export_model(m)))
    w['Assumptions']['D2']=.2;w['Assumptions']['E2']='Human reviewed growth, 2026-10-28'
    raw=BytesIO();w.save(raw)
    out=api().import_assumptions(raw.getvalue(),m)
    assert out['scenarios']['base'][0]['income']['revenue']==1200
    assert out['model_version']!=m['model_version']
    assert out['human_review']=='pending'
    assert out['assumption_changes'][0]['previous']==.1
    assert m['assumptions']['base'][0]['values']['growth']==.1


def test_excel_numeric_roundtrip_without_edits_keeps_model_identity():
    m=fixture();w=load_workbook(BytesIO(api().export_model(m)))
    raw=BytesIO();w.save(raw)  # Native Excel also stores integral doubles as integer literals.
    out=api().import_assumptions(raw.getvalue(),m)
    assert out['assumption_changes']==[]
    assert out['model_version']==m['model_version']


@pytest.mark.parametrize('change',['formula','baseline','identity','source_formula','duplicate','blank_source'])
def test_import_rejects_modified_baseline_stale_identity_or_untrusted_cells(change):
    m=fixture();w=load_workbook(BytesIO(api().export_model(m)))
    if change=='formula':w['Assumptions']['D2']='=1+1'
    if change=='baseline':w['Baseline']['B2']=9
    if change=='identity':w['Review']['B3']='stale-model'
    if change=='source_formula':w['Assumptions']['E2']='=HYPERLINK("https://evil.invalid")'
    if change=='duplicate':w['Assumptions'].append(['base',2026,'growth',.2,'human'])
    if change=='blank_source':w['Assumptions']['E2']=''
    raw=BytesIO();w.save(raw)
    with pytest.raises(ValueError):api().import_assumptions(raw.getvalue(),m)


def test_source_strings_never_become_formulas_or_clickable_remote_links():
    b=baseline();b['source_locator']='=HYPERLINK("https://evil.invalid")'
    m=project(b,scenarios(),as_of='2026-10-27')
    w=load_workbook(BytesIO(api().export_model(m)))
    assert w['Baseline']['C2'].data_type=='s'
    assert w['Baseline']['C2'].hyperlink is None


def test_pending_model_survives_sorted_json_checkpoint_roundtrip():
    import json
    s=scenarios();s['bull'][0]['values']['growth']=None;s['bear'][0]['values']['growth']=None
    m=project(baseline(),s,as_of='2026-10-27')
    saved=json.loads(json.dumps(m,sort_keys=True))
    raw=api().export_model(saved)
    assert api().import_assumptions(raw,saved)['status']=='pending_assumptions'


def test_editable_workbook_guards_forecasts_after_required_values_or_sources_are_cleared():
    complete=load_workbook(BytesIO(api().export_model(fixture())))
    pending_inputs=scenarios();pending_inputs['base'][0]['values']['growth']=None
    pending=load_workbook(BytesIO(api().export_model(project(baseline(),pending_inputs,as_of='2026-10-27'))))
    # A shared prerequisite checks both value and source on every assumption row.
    assert complete['Assumptions']['G2'].data_type=='f'
    assert complete['Assumptions']['F2'].value=='=IF(ISNUMBER(D2),LEN(TRIM(E2))>0,FALSE)'
    for book in (complete,pending):
        for tab in ('Base','Bull','Bear'):
            for cell in ('C5','C12','C17','C38','C44'):
                assert book[tab][cell].value.startswith('=IF(Assumptions!$G$2,')
    cached=load_workbook(BytesIO(api().export_model(project(baseline(),pending_inputs,as_of='2026-10-27'))),data_only=True)
    assert cached['Assumptions']['G2'].value is False
    assert cached['Base']['C44'].value is None


def test_every_statement_formula_matches_cache_across_distinct_cases_and_years():
    from openpyxl.utils.cell import range_boundaries,get_column_letter
    s=scenarios()
    for name,growth in [('base',.1),('bull',.3),('bear',-.2)]:
        s[name][0]['values']['growth']=growth
        s[name].append(dict(deepcopy(s[name][0]),year=2027))
    data=api().export_model(project(baseline(),s,as_of='2026-10-27'))
    formulas=load_workbook(BytesIO(data));cached=load_workbook(BytesIO(data),data_only=True)
    memo={}
    def value(sheet,cell):
        key=(sheet,cell)
        if key in memo:return memo[key]
        raw=formulas[sheet][cell].value
        if not isinstance(raw,str) or not raw.startswith('='):return raw
        raw=raw.replace('$','')
        prefix='=IF(Assumptions!G2,'
        if raw.startswith(prefix):
            if not value('Assumptions','G2'):return ''
            raw='='+raw[len(prefix):-4]
        def substitute(match):
            tab=match[1] or sheet
            if ':' in match[2]:
                c1,r1,c2,r2=range_boundaries(match[2])
                return repr([value(tab,f'{get_column_letter(c)}{r}') for r in range(r1,r2+1) for c in range(c1,c2+1)])
            return repr(value(tab,match[2]))
        expression=re.sub(r'(?:(\w+)!)?([A-Z]+[0-9]+(?::[A-Z]+[0-9]+)?)',substitute,raw[1:])
        expression=re.sub(r'(?<![<>!])=(?!=)','==',expression).replace('TRUE','True').replace('FALSE','False')
        result=eval(expression,{'__builtins__':{}},{'SUM':sum,'MAX':max,'IF':lambda c,a,b:a if c else b,
            'COUNTIF':lambda rows,v:sum(r==v for r in rows),'ISNUMBER':lambda v:isinstance(v,(int,float)) and not isinstance(v,bool),
            'LEN':len,'TRIM':lambda v:str(v or '').strip()})
        memo[key]=result;return result
    for tab in ('Base','Bull','Bear'):
        for row in formulas[tab]:
            for cell in row:
                if cell.data_type=='f':
                    assert value(tab,cell.coordinate)==pytest.approx(cached[tab][cell.coordinate].value)
    assert cached['Bull']['D5'].value==pytest.approx(1690)
    assert cached['Bear']['D5'].value==pytest.approx(640)
    formulas['Assumptions']['D2']=None;memo.clear()
    for tab in ('Base','Bull','Bear'):assert value(tab,'C44')==''
    formulas['Assumptions']['D2']=.1;formulas['Assumptions']['E2']='   ';memo.clear()
    assert value('Base','C5')==''
    formulas['Assumptions']['D2']=0.;formulas['Assumptions']['E2']='explicit zero growth';memo.clear()
    assert value('Base','C5')==1000
