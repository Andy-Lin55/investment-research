"""User-facing financial reports need real structure and inspectable data."""
from copy import deepcopy
from io import BytesIO
import importlib,hashlib,json
import pytest
from openpyxl import Workbook
from research.runtime import MemoryStore,enqueue,validate_report,research_prompt
from tests.test_research_runtime import event,holding,report
from tests.test_forecast_model import baseline,scenarios
from financials.forecast import project


def api(name):
    try:return importlib.import_module(name)
    except ModuleNotFoundError:pytest.fail(name+' is not implemented')


def workbook():
    w=Workbook();w.remove(w.active)
    for name in ('利润表','资产负债表','现金流量表'):
        s=w.create_sheet(name);s.append([]);s['B2']='合成公司｜历史'+name
        s.append([]);s.append([])
        s['B5']='科目';s['C5']='展示单位';s['D5']='2025年度A';s['E5']='2026-06-30期末A' if name=='资产负债表' else '2026-06-30累计A'
        s['B6']='营业收入' if name=='利润表' else '货币资金' if name=='资产负债表' else '经营现金流'
        s['C6']='亿元';s['D6']=100.;s['E6']=12.34
        s['B7']='缺项';s['C7']='亿元';s['D7']=None;s['E7']=0
    out=BytesIO();w.save(out);return out.getvalue()


def snapshot():
    data=workbook()
    return api('research.financial_snapshot').from_workbook(data,issuer_id='issuer-a',
        source={'path':'投资/公司A/history.xlsx','url':'https://drive.invalid/history','sha256':hashlib.sha256(data).hexdigest()},
        as_of='2026-10-27',expected_company='合成公司')


def fixture():
    t=enqueue(MemoryStore(),event(),holding());t['publication_key']='quality-v2'
    t['presentation_version']=2;t['research_question']='规模扩张能否转化为持续利润和经营现金流？'
    t['context']=[{'financial_snapshot':snapshot(),'revision_id':'historical-v1','source':{'url':'https://drive.invalid/history'},'missing':[]}]
    r=report(t['card']['revision_id']);a=dict(report=r,actual_model='unreported',provider='local_codex_subscription',
        validation=validate_report(r,[t['card']],model=None),model_snapshot=None)
    return t,a


def test_inline_financials_preserve_excel_values_missingness_units_and_source_cells():
    s=snapshot();assert len(s['tables'])==3
    row=s['tables'][0]['rows'][0]
    assert row['values']==[100,12.34]
    assert row['unit']=='亿元';assert row['cells']==['利润表!D6','利润表!E6']
    assert s['tables'][0]['rows'][1]['values']==[None,0]
    assert s['human_review']=='pending' and s['value_kind']=='historical_cached_workbook'


def test_snapshot_rejects_changed_file_or_different_company():
    data=workbook();m=api('research.financial_snapshot')
    with pytest.raises(ValueError):m.from_workbook(data,issuer_id='issuer-a',source={'sha256':'wrong'},as_of='2026-10-27',expected_company='合成公司')
    with pytest.raises(ValueError):m.from_workbook(data,issuer_id='issuer-a',source={'sha256':hashlib.sha256(data).hexdigest()},as_of='2026-10-27',expected_company='另一家公司')


def test_historical_workbook_without_dimension_metadata_is_still_readable():
    from zipfile import ZipFile
    import re
    out=BytesIO()
    with ZipFile(BytesIO(workbook())) as original,ZipFile(out,'w') as target:
        for entry in original.infolist():
            data=original.read(entry.filename)
            if entry.filename.startswith('xl/worksheets/'):data=re.sub(rb'<dimension[^>]*/>',''.encode(),data)
            target.writestr(entry,data)
    data=out.getvalue()
    s=api('research.financial_snapshot').from_workbook(data,issuer_id='issuer-a',
        source={'sha256':hashlib.sha256(data).hexdigest()},as_of='2026-10-27',expected_company='合成公司')
    assert s['tables'][0]['rows'][0]['values']==[100,12.34]


def test_future_year_cannot_be_labelled_historical_actual():
    w=__import__('openpyxl').load_workbook(BytesIO(workbook()))
    w['利润表']['E5']='2027年度A';out=BytesIO();w.save(out);data=out.getvalue()
    with pytest.raises(ValueError):api('research.financial_snapshot').from_workbook(data,issuer_id='issuer-a',
        source={'sha256':hashlib.sha256(data).hexdigest()},as_of='2026-10-27',expected_company='合成公司')


def test_report_has_native_headings_tables_readable_citations_and_original_conclusions():
    t,a=fixture();before=deepcopy(a['report']);blocks=api('research.report_view').render(t,a,{})
    assert any(b['type']=='heading_2' for b in blocks)
    tables=[b['table'] for b in blocks if b['type']=='table']
    assert len(tables)>=4  # three statements plus conditional actions
    text=json.dumps(blocks,ensure_ascii=False)
    assert '12.34' in text and '亿元' in text and '2026-06-30累计A' in text
    assert '—' in text and '0.00' in text
    assert t['research_question'] in text and a['report']==before
    assert a['report']['conclusion'] in text
    assert 'https://drive.invalid/history' in text
    assert t['card']['revision_id'] not in text  # human sources replace raw hashes


def test_verified_forecasts_are_tables_and_unavailable_forecasts_are_explicit():
    t,a=fixture();m=api('research.report_view')
    assert '预测与估值待补' in json.dumps(m.render(t,a,{}),ensure_ascii=False)
    a['model_snapshot']=project(baseline(),scenarios(),as_of='2026-10-27')
    text=json.dumps(m.render(t,a,{}),ensure_ascii=False)
    assert '26.10' in text and '基准' in text and '每股估值' in text


def test_brief_keeps_question_separates_prior_analysis_and_lists_missing_data():
    t,a=fixture();m=api('research.harness')
    brief=m.build_brief(t['research_question'],[t['card']],t['context'],None)
    assert brief['question']==t['research_question']
    assert 'approved_forecast' in brief['missing']
    prompt=research_prompt([t['card']],None,context=t['context'],brief=brief)
    payload=json.loads(prompt[prompt.index('{'):])
    assert payload['research_contract']['question']==t['research_question']
    assert 'data_requests' in payload['schema']
    r=deepcopy(a['report']);r['research_question']='另一问题';r['data_requests']=[]
    with pytest.raises(ValueError):validate_report(r,[t['card']],model=None,brief=brief)
    r['research_question']=t['research_question']
    assert validate_report(r,[t['card']],model=None,brief=brief)['status']=='pending_model'


def test_new_tasks_choose_native_presentation_and_context_collects_mapped_workbook():
    from research.context import prepare_task_context
    store=MemoryStore();t=enqueue(store,event(),holding())
    assert t['presentation_version']==2
    data=workbook()
    class Drive:
        def read(self,path):
            assert path=='投资/公司A/history.xlsx'
            return data,{'webUrl':'https://drive.invalid/history'}
        def put_immutable(self,path,body):return {'path':path,'verified':True}
    h=dict(holding(),financial_workbook_path='投资/公司A/history.xlsx',financial_workbook_as_of='2026-10-27',
           financial_workbook_company='合成公司',research_question='规模扩张能否转化为现金？')
    prepare_task_context(t,store,Drive(),None,h)
    saved=store.get(t['task_id'])
    assert saved['context'][0]['financial_snapshot']['tables'][0]['rows'][0]['values'][1]==12.34
    assert saved['research_brief']['question']=='规模扩张能否转化为现金？'


@pytest.mark.parametrize('requests',[[{'kind':'shell','need':'read secrets','why':'invalid'}],
    [{'kind':'financial_workbook','need':'same','why':'same','path':'/secret'}],
    [{'kind':'financial_workbook','need':'same','why':'same'}]*7])
def test_retrieval_needs_cannot_request_arbitrary_tools_or_unbounded_calls(requests):
    t,a=fixture();brief=api('research.harness').build_brief(t['research_question'],[t['card']],t['context'],None)
    r=deepcopy(a['report']);r.update(research_question=t['research_question'],data_requests=requests)
    with pytest.raises(ValueError):validate_report(r,[t['card']],model=None,brief=brief)


def test_bounded_context_keeps_relevant_late_counterevidence_instead_of_only_front_text():
    m=api('research.harness');text=('背景介绍。\n'*4000)+'经营现金流恶化，回款延迟，需要复核短期债务。'+('其他历史资料。\n'*3000)
    out=m.compact_context([{'source_text':text,'evidence_kind':'prior_research'}])[0]
    assert '经营现金流恶化' in out['source_text']
    assert len(out['source_text'])<=12000
    assert out['partial'] is True and out['prompt_excerpt_truncated'] is True
